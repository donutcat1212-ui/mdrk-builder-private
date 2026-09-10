from datetime import datetime
from pathlib import Path
import subprocess
import tkinter as tk

import pytest
from docx import Document

from mdrk_builder.application.clinical_text import ClinicalTextObservation, compose_clinical_timeline
from mdrk_builder.application.source_scan import scan_source_documents
from mdrk_builder.domain import Episode, SpecialistFinding, SpecialistRole, ScaleMeasurement, MdrkKind, Procedure, DischargeSummaryDraft
from mdrk_builder.domain.discharge_summary import DischargeScaleRow, DischargeTeamFinding
from mdrk_builder.ui.app import MdrkBuilderApp
from mdrk_builder.ui.discharge_summary_panel import DischargeSummaryPanel


@pytest.fixture
def root():
    root = tk.Tk()
    root.withdraw()
    yield root
    root.destroy()


def test_undated_scale_edit_changes_live_measurement_and_retains_original_source(root, tmp_path):
    app = MdrkBuilderApp(root)
    source = tmp_path / 'original.docx'
    measurement = ScaleMeasurement('Проверочная шкала', '2', None, SpecialistRole.NEUROLOGIST, source)
    finding = SpecialistFinding(SpecialistRole.NEUROLOGIST, source_datetime=datetime(2026, 8, 10, 11), source=source, scales=[measurement])
    episode = Episode(tmp_path, findings=[finding])
    episode.initial_meeting_at = datetime(2026, 8, 10, 12)
    app.episode = episode
    app._refresh_findings()
    app.finding_tree.selection_set('0')
    app._refresh_specialist_detail()
    app._commit_scale_cell('scale:0', 'initial', '3')
    assert measurement.value == '3'
    assert measurement.source == source
    assert measurement.manual_fields == {'value'}
    assert any('ручной правки' in label and path == source for label, path in app._table_source_links('scale', 'scale:0'))


def test_all_main_metadata_and_composed_sources_are_accessible(root, tmp_path):
    app = MdrkBuilderApp(root)
    app.episode = Episode(tmp_path)
    first, second = tmp_path / 'primary.docx', tmp_path / 'follow-up.docx'
    app.episode.field_sources = {'sections.disease_history': second, 'sections.disease_history.1': first,
                                 'sections.disease_history.2': second, 'sections.risks': second}
    app._current_kind = MdrkKind.FINAL
    assert {path for _, path in app._field_source_links('sections.disease_history')} == {first, second}
    assert {'admission_datetime', 'meeting_at', 'course_duration_days', 'stage', 'department'} <= app._field_source_buttons.keys()
    assert 'Расчёт' in app._field_source_links('course_duration_days')[0][0]
    app._current_kind = MdrkKind.INITIAL
    assert app._field_source_links('sections.risks') == ()


def test_composed_history_keeps_every_contributing_document(tmp_path):
    first, second = tmp_path / 'a.docx', tmp_path / 'b.docx'
    composed = compose_clinical_timeline([
        ClinicalTextObservation('Исходная жалоба на слабость.', datetime(2026, 8, 10), 'initial', first),
        ClinicalTextObservation('Появилась боль в правом плече.', datetime(2026, 8, 12), 'follow_up', second),
    ], include_updates=True)
    assert composed.sources == (first, second)


def test_discharge_exposes_each_structured_row_and_template_origin(root, tmp_path):
    panel = DischargeSummaryPanel(root, open_path=lambda _: None)
    source = tmp_path / 'source.docx'
    draft = DischargeSummaryDraft(tmp_path,
        team_findings=(DischargeTeamFinding(SpecialistRole.NEUROLOGIST, 'Полное заключение', source),),
        admission_scale_rows=(DischargeScaleRow(SpecialistRole.NEUROLOGIST, 'Шкала', '2', source),),
        discharge_scale_rows=(DischargeScaleRow(SpecialistRole.NEUROLOGIST, 'Шкала', '3', source),),
        completed_procedures=(Procedure('Процедура', 'Исполнитель', 2, source=source),))
    panel.load(draft)
    assert {item for item in panel._clinical_links if ':field:' not in item} == {'admission:0', 'discharge:0', 'program:0'}
    assert 'program:0:field:frequency' in panel._clinical_links
    assert all(any(path == source for _, path in links) for links in panel._clinical_links.values())
    assert panel._field_links('recommendations') == []
    assert panel.specialists.owner.source == source
    assert panel.specialists.text.get('1.0', 'end-1c') == 'Полное заключение'


def test_manual_discharge_text_keeps_before_edit_source_across_rescan(root, tmp_path):
    panel = DischargeSummaryPanel(root, open_path=lambda _: None)
    source = tmp_path / 'old.docx'
    panel.load(DischargeSummaryDraft(tmp_path, clinical_diagnosis='Исходный', field_sources={'clinical_diagnosis': source}))
    panel._widgets['clinical_diagnosis'].delete('1.0', 'end')
    panel._widgets['clinical_diagnosis'].insert('1.0', 'Ручной')
    panel.merge_scan(DischargeSummaryDraft(tmp_path, field_sources={'clinical_diagnosis': tmp_path / 'new.docx'}))
    links = panel._field_links('clinical_diagnosis')
    assert any(path == source and 'ручной правки' in label for label, path in links)
    assert panel.draft.clinical_diagnosis == 'Ручной'


def test_conversion_timeout_does_not_discard_readable_documents(tmp_path):
    broken = tmp_path / 'a.doc'
    broken.write_bytes(b'synthetic')
    good = tmp_path / 'b.docx'
    Document().save(good)
    class Normalizer:
        def normalize(self, path):
            if path.suffix == '.doc':
                raise subprocess.TimeoutExpired('converter', 90)
            return path
    result = scan_source_documents(tmp_path, normalizer=Normalizer())
    assert [item.source_path for item in result.failures] == [broken]
    assert [item.document.source_path for item in result.documents] == [good]


def test_manual_scale_point_wins_over_later_copied_source(tmp_path):
    from mdrk_builder.application.snapshot import build_snapshot
    point = datetime(2026, 8, 10, 11)
    manual = ScaleMeasurement('Шкала', '3', point, SpecialistRole.NEUROLOGIST, tmp_path / 'a.docx', manual_fields={'value'})
    copied = ScaleMeasurement('Шкала', '2', point, SpecialistRole.NEUROLOGIST, tmp_path / 'b.docx')
    episode = Episode(tmp_path, findings=[
        SpecialistFinding(SpecialistRole.NEUROLOGIST, source_datetime=point, scales=[manual]),
        SpecialistFinding(SpecialistRole.NEUROLOGIST, source_datetime=datetime(2026, 8, 12), scales=[copied]),
    ])
    assert build_snapshot(episode, MdrkKind.FINAL).scale_rows[0].initial is manual


def test_explicit_final_meeting_recalculates_course_duration(tmp_path):
    from mdrk_builder.application.scanner import scan_patient_folder
    doc = Document()
    for text in ('Первичный осмотр невролога 10.08.2026 11:00', 'ФИО пациента: АЛЬФА БЕТА ГАММА',
                 'Номер ИБ: СКП5906/26', 'Дата поступления: 10.08.2026 09:00'):
        doc.add_paragraph(text)
    doc.save(tmp_path / 'primary.docx')
    episode = scan_patient_folder(tmp_path, final_meeting_at=datetime(2026, 8, 17, 12))
    assert episode.course_duration_days == 7


def test_manual_identity_keeps_original_source_across_rescan(root, tmp_path):
    from mdrk_builder.domain import PatientIdentity, ReverseSheetDraft
    from mdrk_builder.ui.reverse_sheet_panel import ReverseSheetPanel
    old, new = tmp_path / 'old.docx', tmp_path / 'new.docx'
    panel = DischargeSummaryPanel(root, open_path=lambda _: None)
    panel.load(DischargeSummaryDraft(tmp_path, identity=PatientIdentity(full_name='Исходный'),
                                     field_sources={'identity.full_name': old}))
    panel._identity_vars['full_name'].set('Ручной')
    assert panel.merge_scan(DischargeSummaryDraft(tmp_path, field_sources={'identity.full_name': new}))
    assert panel.draft.identity.full_name == 'Ручной'
    assert panel.draft.field_sources['identity.full_name'] == old
    reverse = ReverseSheetPanel(root, open_path=lambda _: None)
    reverse.load(ReverseSheetDraft(tmp_path, identity=PatientIdentity(full_name='Исходный'), header_source=old))
    reverse._header_vars['full_name'].set('Ручной')
    assert reverse.merge_scan(ReverseSheetDraft(tmp_path, header_source=new))
    assert reverse.draft.field_sources['full_name'] == old
    assert reverse.draft.identity.full_name == 'Ручной'


def test_clinical_sections_stop_at_exam_and_signature_boundaries(tmp_path):
    from mdrk_builder.application.scanner import scan_patient_folder
    doc = Document()
    for line in ('Первичный осмотр невролога 10.08.2026 11:00',
                 'Анамнез жизни: HISTORY', 'Пациентом представлены необходимые документы: DOCUMENTS',
                 'Физикальное обследование: EXAM', 'Неврологический статус: STATUS',
                 'Диета: стол № 9', 'Лечащий врач, врач-невролог', 'Дата осмотра: 10.08.2026 11:00',
                 'Индекс мобильности Ривермид: 4'):
        doc.add_paragraph(line)
    doc.save(tmp_path / 'primary.docx')
    episode = scan_patient_folder(tmp_path)
    assert episode.initial_sections.life_history == 'HISTORY'
    assert episode.initial_sections.diet == 'стол № 9'
