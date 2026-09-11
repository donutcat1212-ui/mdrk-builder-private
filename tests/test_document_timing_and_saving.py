from copy import deepcopy
from datetime import datetime
import hashlib
import tkinter as tk
from tkinter import ttk

import pytest
from docx import Document

from mdrk_builder.application.discharge_summary import scan_discharge_summary
from mdrk_builder.application.scanner import scan_patient_folder
from mdrk_builder.application.snapshot import build_snapshot
from mdrk_builder.domain import DischargeSummaryDraft, MdrkKind, ReverseSheetDraft
from mdrk_builder.domain.document_dates import end_of_day
from mdrk_builder.infrastructure.discharge_summary_writer import write_discharge_summary_docx
from mdrk_builder.infrastructure.docx_writer import write_mdrk_docx
from mdrk_builder.infrastructure.draft_store import encode, decode
from test_daily_workflow_ui import app
from test_discharge_summary import _primary_lines, _discharge_lines, _write_document
from test_document_order import _assessment, SYNTHETIC_NAME
from test_docx_writer import _representative_episode


def test_final_documents_include_later_times_but_exclude_the_next_day(tmp_path):
    _assessment(tmp_path / "primary.docx", (*_primary_lines(full_name=SYNTHETIC_NAME),
        "Дата осмотра: 10.08.2026 11:00"), 3, 4)
    for day, hour, minute, qualifier, score in [(17, 23, 59, 2, 9), (18, 0, 0, 1, 14)]:
        _assessment(tmp_path / f"follow-up-{day}.docx", (
            f"Повторный осмотр невролога {day}.08.2026 {hour:02d}:{minute:02d}",
            f"ФИО пациента: {SYNTHETIC_NAME}", "Номер ИБ: СКП5906/26",
            f"Клинический диагноз: DAY_{day}",), qualifier, score)
    mis = tmp_path / "mis.docx"
    _write_document(mis, _discharge_lines(full_name=SYNTHETIC_NAME))
    hashes = {p: hashlib.sha256(p.read_bytes()).digest() for p in tmp_path.glob("*.docx")}
    at = datetime(2026, 8, 17, 11)
    episode = scan_patient_folder(tmp_path, initial_meeting_at=datetime(2026, 8, 11, 8), final_meeting_at=at)
    assert episode.meeting_at(MdrkKind.FINAL) == at
    assert episode.assessment_at(MdrkKind.FINAL) == end_of_day(at)
    assert any(source.path == mis for source in episode.sources)
    assert "DISCHARGE DIAGNOSIS" in episode.sections.clinical_diagnosis
    assert "DAY_17" in episode.sections.clinical_diagnosis and "DAY_18" not in episode.sections.clinical_diagnosis
    assert "DISCHARGE DIAGNOSIS" not in episode.initial_sections.clinical_diagnosis
    for candidate in (episode, decode(encode(episode))):
        snapshot = build_snapshot(candidate, MdrkKind.FINAL)
        assert next(row for row in snapshot.scale_rows if "Ривермид" in row.name).current.value == "9"
        assert next(row for row in snapshot.icf_domains if row.code == "s110").final.value == 2
    draft = scan_discharge_summary(tmp_path, discharge_datetime_override=datetime(2026, 8, 17, 7))
    assert draft.discharge_datetime == datetime(2026, 8, 17, 12)
    assert next(row for row in draft.discharge_scale_rows if "Ривермид" in row.name).value == "9"
    assert next(row for row in draft.icf_domains if row.code == "s110").final.value == 2
    mdrk = Document(write_mdrk_docx(episode, MdrkKind.FINAL, tmp_path / "mdrk2.docx", ignore_issues=True))
    discharge = Document(write_discharge_summary_docx(draft, tmp_path / "discharge.docx", ignore_issues=True))
    assert any("время: 11 час. 00 мин." in p.text for p in mdrk.paragraphs)
    assert any("17.08.2026 12:00" in p.text for p in discharge.paragraphs)
    assert all(hashlib.sha256(p.read_bytes()).digest() == digest for p, digest in hashes.items())


def test_old_discharge_draft_uses_noon_without_requiring_a_same_day_rescan(tmp_path):
    start, end = datetime(2026, 8, 10, 10), datetime(2026, 8, 17, 18, 30)
    draft = DischargeSummaryDraft(tmp_path, admission_datetime=start, discharge_datetime=end,
        projection_period=(start, end), header_text="Дата и время выписки: 17.08.2026 18:30")
    restored = decode(encode(draft))
    original = deepcopy(restored)
    output = Document(write_discharge_summary_docx(restored, tmp_path / "noon.docx"))
    assert "Дата и время выписки: 17.08.2026 12:00" in [p.text for p in output.paragraphs]
    assert restored == original


@pytest.mark.parametrize("button_text", ["Открыть документ", "Показать в папке"])
def test_success_dialog_closes_when_opening_result(app, monkeypatch, button_text):
    created = app.episode.folder / "result.docx"
    opened = []
    monkeypatch.setattr(app, "_save_workspace", lambda: None)
    monkeypatch.setattr(app, "_open_path", opened.append)
    app._finish_save(created, None)
    window = next(w for w in app.root.winfo_children() if isinstance(w, tk.Toplevel))
    buttons = [w for frame in window.winfo_children() for w in frame.winfo_children() if isinstance(w, ttk.Button)]
    next(b for b in buttons if b.cget("text") == button_text).invoke()
    assert not window.winfo_exists()
    assert opened == [created if button_text == "Открыть документ" else created.parent]


@pytest.mark.parametrize("document", ["mdrk1", "mdrk2", "discharge", "reverse"])
def test_save_dialog_starts_in_the_current_episode_folder(app, monkeypatch, document):
    captured = []
    monkeypatch.setattr("tkinter.filedialog.asksaveasfilename", lambda **kwargs: captured.append(kwargs) or "")
    for module in ("app", "discharge_summary_panel", "reverse_sheet_panel"):
        monkeypatch.setattr(f"mdrk_builder.ui.{module}.confirm_generation_with_issues", lambda *a, **k: True)
    episode = _representative_episode(app.episode.folder)
    app.episode, app._scan_baseline = episode, deepcopy(episode)
    app._current_kind = MdrkKind.FINAL if document == "mdrk2" else MdrkKind.INITIAL
    app.document_var.set(document)
    app._populate_from_episode()
    if document == "discharge":
        app.discharge_workspace.load(DischargeSummaryDraft(episode.folder))
        app.discharge_workspace.save()
    elif document == "reverse":
        app.reverse_workspace.load(ReverseSheetDraft(episode.folder))
        app.reverse_workspace.save()
    else:
        app._generate()
    assert len(captured) == 1
    assert captured[0]["initialdir"] == str(episode.folder)
