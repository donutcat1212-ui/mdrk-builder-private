"""Source-folder to saved draft to rescan/export release acceptance scenarios."""
from copy import deepcopy
from datetime import datetime
import hashlib
import tkinter as tk

from docx import Document
import pytest

from mdrk_builder.application.discharge_summary import scan_discharge_summary
from mdrk_builder.application.scanner import scan_patient_folder
from mdrk_builder.application.reverse_sheet import scan_reverse_sheet
from mdrk_builder.application.snapshot import build_snapshot
from mdrk_builder.domain import MdrkKind
from mdrk_builder.infrastructure.draft_store import save_draft, load_draft
from mdrk_builder.infrastructure.discharge_summary_writer import write_discharge_summary_docx
from mdrk_builder.infrastructure.docx_writer import write_mdrk_docx
from mdrk_builder.infrastructure.reverse_sheet_writer import write_reverse_sheet_docx
from mdrk_builder.ui.document_panels import DischargeSummaryPanel
from test_discharge_summary import _write_document, _primary_lines, _discharge_lines


def assignment(path, days):
    document = Document()
    document.add_paragraph('Лист назначений')
    document.add_paragraph('ФИО пациента: АЛЬФА БЕТА ГАММА')
    document.add_paragraph('Номер ИБ: СКП5906/26')
    table = document.add_table(rows=2, cols=len(days) + 3)
    for cell, value in zip(table.rows[0].cells, ['Назначения', 'Время', 'Длительность', *days]):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ['A19.23.001 ЛФК', '10:00', '30 мин', *['+'] * len(days)]):
        cell.text = value
    document.save(path)


@pytest.mark.parametrize('scenario', ['missing', 'continuation', 'late_update'])
def test_folder_edit_rescan_draft_and_all_exports(tmp_path, scenario):
    folder = tmp_path / scenario
    folder.mkdir()
    _write_document(folder / 'первичный осмотр невролога.docx', _primary_lines(full_name='АЛЬФА БЕТА ГАММА'))
    _write_document(folder / 'выписной эпикриз.docx', _discharge_lines(full_name='АЛЬФА БЕТА ГАММА'))
    if scenario != 'missing':
        assignment(folder / 'лист назначений.docx', ['11.08.2026', '12.08.2026'])
        assignment(folder / 'лист назначений продолжение.docx', ['12.08.2026', '16.08.2026', '19.08.2026'])
    if scenario == 'late_update':
        _write_document(folder / 'заключительный осмотр невролога.docx', (
            'Заключительный осмотр невролога 16.08.2026 10:00',
            'ФИО пациента: АЛЬФА БЕТА ГАММА', 'Номер ИБ: СКП5906/26',
            'Двигательный режим: свободный', 'Состояние при выписке: УЛУЧШЕНИЕ',
            'Рекомендации: ИТОГОВЫЕ РЕКОМЕНДАЦИИ'))
    original = {p: hashlib.sha256(p.read_bytes()).digest() for p in folder.glob('*.docx')}
    draft = scan_discharge_summary(folder)
    root = tk.Tk()
    root.withdraw()
    try:
        panel = DischargeSummaryPanel(root, open_path=lambda path: None)
        panel.load(draft)
        panel._widgets['recommendations'].insert('1.0', 'РУЧНАЯ ПРАВКА\n')
        assert panel.apply()
        assert panel.merge_scan(scan_discharge_summary(folder))
        assert panel.draft.recommendations.startswith('РУЧНАЯ ПРАВКА')
        save_draft(tmp_path / 'draft.json', panel.draft)
        restored = load_draft(tmp_path / 'draft.json')
        assert restored.recommendations == panel.draft.recommendations
    finally:
        root.destroy()
    if scenario != 'missing':
        assert [p.actual_count for p in restored.completed_procedures] == [3]
    if scenario == 'late_update':
        assert restored.movement_regimen == 'свободный'
    output = tmp_path / 'output'
    output.mkdir()
    episode = scan_patient_folder(folder, initial_meeting_at=datetime(2026, 8, 11, 8), final_meeting_at=datetime(2026, 8, 15, 12))
    if scenario != 'missing':
        assert build_snapshot(episode, MdrkKind.FINAL).procedures[0].actual_count == 2
    paths = [write_mdrk_docx(episode, kind, output / (kind.value + '.docx')) for kind in MdrkKind]
    paths.append(write_discharge_summary_docx(restored, output / 'выписка.docx'))
    reverse = scan_reverse_sheet(folder)
    paths.append(write_reverse_sheet_docx(reverse, output / 'оборотный лист.docx'))
    assert len(paths) == 4
    for path in paths:
        document = Document(path)
        assert document.paragraphs or document.tables
    assert all(hashlib.sha256(p.read_bytes()).digest() == digest for p, digest in original.items())
