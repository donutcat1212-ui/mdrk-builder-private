from copy import deepcopy
from datetime import datetime
import tkinter as tk

import pytest

from mdrk_builder.domain import Episode, DischargeSummaryDraft, ReverseSheetDraft, MdrkKind, Procedure
from mdrk_builder.ui.app import MdrkBuilderApp
from mdrk_builder.ui.document_panels import DischargeSummaryPanel


@pytest.fixture
def app(tmp_path, monkeypatch):
    root=tk.Tk();root.withdraw()
    errors=[]
    monkeypatch.setattr('tkinter.messagebox.showerror',lambda *a,**k:errors.append(a))
    monkeypatch.setattr('tkinter.messagebox.askyesno',lambda *a,**k:True)
    ui=MdrkBuilderApp(root)
    ui._set_folder_field(str(tmp_path))
    ui.episode=Episode(tmp_path, admission_datetime=datetime(2026,8,10), final_meeting_at=datetime(2026,8,17), course_duration_days=7)
    ui._scan_baseline=deepcopy(ui.episode)
    ui._populate_from_episode()
    ui._test_errors=errors
    yield ui
    root.destroy()


def test_automatic_days_and_manual_override_survive_date_edits(app):
    app._entry_variables['admission'].set('12.08.2026 00:00')
    app.root.update()
    assert app._entry_variables['duration'].get()=='5'
    app._entry_variables['duration'].set('9')
    app._entry_variables['admission'].set('13.08.2026 00:00')
    app.root.update()
    assert app._entry_variables['duration'].get()=='9'
    app._automatic_duration()
    assert app._entry_variables['duration'].get()=='4'


def test_header_and_warning_refresh_without_rescan(app):
    panel=app.discharge_workspace
    panel.load(DischargeSummaryDraft(app.episode.folder,header_text='ФИО: Старое',discharge_datetime=datetime(2026,8,18)))
    panel._identity_vars['full_name'].set('Новое имя')
    panel._widgets['clinical_diagnosis'].insert('1.0','Диагноз')
    app.root.update()
    assert 'Новое имя' in panel._widgets['header_text'].get('1.0','end-1c')
    assert not any(i.field=='clinical_diagnosis' for i in panel.draft.issues)
    panel._widgets['clinical_diagnosis'].delete('1.0','end')
    app.root.update()
    assert any(i.field=='clinical_diagnosis' for i in panel.draft.issues)


def test_patient_change_clears_text_undo(app):
    panel=app.discharge_workspace
    panel.load(DischargeSummaryDraft(app.episode.folder,clinical_diagnosis='Пациент А'))
    panel._widgets['clinical_diagnosis'].insert('end',' правка')
    panel.load(DischargeSummaryDraft(app.episode.folder,clinical_diagnosis='Пациент Б'))
    with pytest.raises(tk.TclError):panel._widgets['clinical_diagnosis'].edit_undo()
    assert panel._widgets['clinical_diagnosis'].get('1.0','end-1c')=='Пациент Б'


def test_workspace_recovers_unsaved_intermediate_invalid_input(app):
    app.discharge_workspace.load(DischargeSummaryDraft(app.episode.folder,clinical_diagnosis='Текст'))
    app.discharge_draft=app.discharge_workspace.draft
    app.document_var.set('discharge');app._previous_document='discharge'
    app.discharge_workspace._identity_vars['birth_date'].set('12.')
    app.discharge_workspace._widgets['clinical_diagnosis'].insert('end',' правка')
    assert app._save_workspace()
    app.discharge_workspace._identity_vars['birth_date'].set('')
    assert app._restore_workspace(app.episode.folder)
    assert app.discharge_workspace._identity_vars['birth_date'].get()=='12.'
    assert app.discharge_workspace._widgets['clinical_diagnosis'].get('1.0','end-1c')=='Текст правка'
    assert not app._test_errors


def test_identity_moves_between_existing_document_tabs(app):
    app.reverse_workspace.load(ReverseSheetDraft(app.episode.folder))
    app.reverse_draft=app.reverse_workspace.draft
    app.discharge_workspace.load(DischargeSummaryDraft(app.episode.folder))
    app.discharge_draft=app.discharge_workspace.draft
    app._entry_variables['full_name'].set('Общее имя')
    app.document_var.set('reverse');app._on_document_changed()
    assert app.reverse_workspace._header_vars['full_name'].get()=='Общее имя'
    app.reverse_workspace._header_vars['full_name'].set('Уточнённое имя')
    app.document_var.set('discharge');app._on_document_changed()
    assert app.discharge_workspace._identity_vars['full_name'].get()=='Уточнённое имя'
    assert app.episode.identity.full_name=='Уточнённое имя'


def test_rescan_preserves_edit_and_new_rows_and_resets_table_history(app):
    panel=app.discharge_workspace
    draft=DischargeSummaryDraft(app.episode.folder,completed_procedures=(Procedure('А','ФТ',3),))
    panel.load(draft)
    panel._table_history.wrap(lambda:setattr(panel.draft.completed_procedures[0],'actual_count',5))()
    fresh=DischargeSummaryDraft(app.episode.folder,completed_procedures=(Procedure('А','ФТ',4),Procedure('Б','ФТ',2)))
    assert panel.merge_scan(fresh)
    assert [r.actual_count for r in panel.draft.completed_procedures]==[5,2]
    panel._table_history.undo()
    assert len(panel.draft.completed_procedures)==2


def test_tab_commits_and_advances_and_shift_tab_returns(app):
    from tkinter import ttk
    from mdrk_builder.ui.inline_tree import InlineTreeEditor
    top=tk.Toplevel(app.root)
    tree=ttk.Treeview(top,columns=('a','b'),show='headings');tree.pack()
    tree.insert('','end',iid='0',values=('one','two'))
    top.update()
    editor=InlineTreeEditor(tree,editable_columns={'a','b'},commit=lambda row,col,value:tree.set(row,col,value))
    editor.edit('0','a');editor._widget.delete(0,'end');editor._widget.insert(0,'changed')
    editor._next_cell('0','a',1)
    assert tree.set('0','a')=='changed'
    assert editor._widget.get()=='two'
    editor._next_cell('0','b',-1)
    assert editor._widget.get()=='changed'
    top.destroy()


def test_failed_rescan_restores_previous_work(app):
    from pathlib import Path
    app._entry_variables['full_name'].set('Ручное имя')
    app._pending_manual_state=app._capture_manual_state()
    app._active_job_folder=app.episode.folder
    app.episode=None
    app._finish_scan(None,RuntimeError('test failure'))
    assert app.episode.identity.full_name=='Ручное имя'
    assert app._entry_variables['full_name'].get()=='Ручное имя'


def test_saved_indicator_follows_changes(app):
    app._refresh_draft_indicator()
    assert 'Черновик не сохранён' in app.root.title()
    app._save_workspace()
    app._refresh_draft_indicator()
    assert 'Черновик не сохранён' not in app.root.title()
    app._entry_variables['full_name'].set('Правка')
    app._refresh_draft_indicator()
    assert 'Черновик не сохранён' in app.root.title()


def test_table_undo_preserves_text_and_identity(app):
    panel=app.discharge_workspace
    panel.load(DischargeSummaryDraft(app.episode.folder,completed_procedures=(Procedure('А','ФТ',3),)))
    panel._widgets['clinical_diagnosis'].insert('1.0','Сохранить текст')
    panel._identity_vars['full_name'].set('Сохранить имя')
    panel._table_history.wrap(lambda:setattr(panel.draft.completed_procedures[0],'actual_count',5))()
    panel._table_history.undo()
    assert panel.draft.completed_procedures[0].actual_count==3
    assert panel._widgets['clinical_diagnosis'].get('1.0','end-1c')=='Сохранить текст'
    assert panel._identity_vars['full_name'].get()=='Сохранить имя'


def test_folder_change_clears_auxiliary_undo_before_scan(app,tmp_path):
    panel=app.discharge_workspace
    panel.load(DischargeSummaryDraft(tmp_path,completed_procedures=(Procedure('А','ФТ',3),)))
    panel._table_history.wrap(lambda:setattr(panel.draft.completed_procedures[0],'actual_count',5))()
    new=tmp_path/'new';new.mkdir()
    app.folder_var.set(str(new))
    panel._table_history.undo()
    assert panel.draft is None
    assert not panel._table_history.past


def test_added_specialist_scale_uses_parent_role_and_both_views(app,monkeypatch):
    from mdrk_builder.domain import DischargeTeamFinding, SpecialistRole
    from types import SimpleNamespace
    panel=app.discharge_workspace
    role=SpecialistRole.PHYSICAL_THERAPIST
    panel.load(DischargeSummaryDraft(app.episode.folder,team_findings=(DischargeTeamFinding(role,''),)))
    panel.clinical_tree.selection_set('team:0')
    monkeypatch.setattr('mdrk_builder.ui.discharge_tables.FieldsDialog',lambda *a:SimpleNamespace(result={
        'name':'Берг','value':'35','initial_value':'20','initial_at':'10.08.2026 00:00','current_at':'18.08.2026 00:00'}))
    panel._edit_clinical_row('add_scale')
    row=panel.draft.team_findings[0].scales[0]
    assert row.role is role
    assert panel.draft.admission_scale_rows[0].value=='20'
    assert panel.draft.discharge_scale_rows[0].value=='35'


def test_source_preview_shows_table_coordinates(app,tmp_path):
    from docx import Document
    from mdrk_builder.application.source_scan import scan_source_documents
    from mdrk_builder.ui.source_preview import preview_source
    path=tmp_path/'source.docx'
    document=Document();document.add_paragraph('Источник');table=document.add_table(rows=1,cols=2)
    table.cell(0,0).text='Берг';table.cell(0,1).text='30';document.save(path)
    app._scan_session.begin(tmp_path)
    scan_source_documents(tmp_path,session=app._scan_session)
    preview_source(app.root,path,'Берг',lambda p:None)
    window=next(child for child in app.root.winfo_children() if isinstance(child,tk.Toplevel))
    text=next(child for child in window.winfo_children() if isinstance(child,tk.Text))
    assert 'Таблица 1, строка 1: Берг | 30' in text.get('1.0','end-1c')
    assert text.tag_ranges('match')


def test_plan_count_edit_does_not_replace_completed_count(app):
    from mdrk_builder.domain import Procedure, MdrkKind
    app.episode.procedures = [Procedure('ЛФК', 'ФТ', 4, planned_count=10)]
    app.kind_var.set(MdrkKind.INITIAL.value)
    app._commit_procedure_cell('0', 'count', '12')
    assert app.episode.procedures[0].planned_count == 12
    assert app.episode.procedures[0].actual_count == 4
    assert str(app.procedure_tree.set('0', 'count')) == '12'
    app.kind_var.set(MdrkKind.FINAL.value)
    app._refresh_procedures()
    assert str(app.procedure_tree.set('0', 'count')) == '4'
