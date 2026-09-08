import tkinter as tk

import pytest
from docx import Document

from mdrk_builder.domain import DischargeSummaryDraft, IcfDomain, IcfQualifier, SpecialistRole
from mdrk_builder.infrastructure.discharge_summary_writer import write_discharge_summary_docx
from mdrk_builder.ui.discharge_summary_panel import DischargeSummaryPanel


def test_icf_panel_matches_docx_and_refreshes_on_rescan(tmp_path):
    try:
        root = tk.Tk()
    except tk.TclError as exc:
        pytest.skip(f"Tk display unavailable: {exc}")
    root.withdraw()
    try:
        opened = []
        panel = DischargeSummaryPanel(root, open_path=opened.append)
        draft = DischargeSummaryDraft(folder=tmp_path)
        draft.final_mdrk_source = tmp_path / "final.docx"
        draft.icf_domains = (
            IcfDomain("d450", "Ходьба", SpecialistRole.PHYSICAL_THERAPIST,
                      initial=IcfQualifier(2), final=IcfQualifier(1),
                      initial_source=tmp_path / "initial.docx", final_source=draft.final_mdrk_source),
            IcfDomain("d640", "Ведение домашнего хозяйства", SpecialistRole.OTHER,
                      initial=IcfQualifier(3)),
            IcfDomain("Pf", "Личные факторы", SpecialistRole.OTHER),
        )
        panel.load(draft)
        assert panel.notebook.tab(1, "text") == "МКФ"
        assert tuple(map(str, panel.icf_tree.item("domain:0", "values")))[2:4] == ("2", "1")
        assert tuple(map(str, panel.icf_tree.item("domain:1", "values")))[2:4] == ("3", "")
        assert tuple(panel.icf_tree.item("domain:2", "values"))[2:] == ("", "", "", "")
        panel.icf_source_button.invoke()
        assert opened == [draft.final_mdrk_source]
        panel.icf_tree.selection_set("domain:0")
        panel._icf_sources._fill_menu()
        panel._icf_sources.menu.invoke(0)
        panel._icf_sources.menu.invoke(1)
        assert opened[-2:] == [tmp_path / "initial.docx", draft.final_mdrk_source]
        output = write_discharge_summary_docx(draft, tmp_path / "output.docx")
        table = next(t for t in Document(output).tables if t.cell(0, 0).text == "МКФ категориальный профиль")
        for index, code in enumerate(("d450", "d640")):
            row = next(r for r in table.rows if r.cells[0].text == code)
            values = tuple(map(str, panel.icf_tree.item(f"domain:{index}", "values")))
            assert values[2:4] == (row.cells[11].text, row.cells[12].text)
        fresh = DischargeSummaryDraft(folder=tmp_path)
        panel.merge_scan(fresh)
        assert not panel.icf_tree.get_children()
        assert panel.icf_source_button.instate(["disabled"])
        assert "не найден" in panel.icf_status.get()
        assert "МДРК-2" not in panel.icf_status.get()
        assert not panel.icf_source_button.winfo_manager()
        fresh.final_mdrk_source = draft.final_mdrk_source
        panel.load(fresh)
        assert "не найдены оценки МКФ" in panel.icf_status.get()
    finally:
        root.destroy()
