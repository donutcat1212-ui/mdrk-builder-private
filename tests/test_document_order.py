"""Both final documents are projections of source facts, in either export order."""
from datetime import datetime
import hashlib

from docx import Document
import pytest

from mdrk_builder.application.discharge_summary import scan_discharge_summary
from mdrk_builder.application.scanner import scan_patient_folder
from mdrk_builder.domain import MdrkKind
from mdrk_builder.infrastructure.discharge_summary_writer import write_discharge_summary_docx
from mdrk_builder.infrastructure.docx_writer import write_mdrk_docx
from test_discharge_summary import _primary_lines, _discharge_lines, _write_document


def _assessment(path, lines, qualifier, scale):
    document = Document()
    for line in (*lines, f"Индекс мобильности Ривермид: {scale}"):
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=15)
    table.cell(0, 0).text = "МКФ"
    table.cell(0, 13).text = "Ответственный специалист"
    table.cell(1, 0).text = "s110"
    table.cell(1, 1).text = "Структура головного мозга"
    table.cell(1, 11).text = str(qualifier)
    document.save(path)


@pytest.mark.parametrize("first", ["mdrk2", "discharge"])
def test_export_order_and_duplicate_outputs_do_not_change_either_document(tmp_path, first):
    _assessment(tmp_path / "primary.docx", (*_primary_lines(),
        "Дата осмотра: 10.08.2026 11:00"), 3, 4)
    _assessment(tmp_path / "follow-up.docx", (
        "Повторный осмотр невролога 16.08.2026 10:00",
        "ФИО пациента: Пациент Тестовый Пример", "Номер ИБ: СКП5906/26",
    ), 2, 9)
    _write_document(tmp_path / "mis.docx", _discharge_lines())
    source_hashes = {path: hashlib.sha256(path.read_bytes()).digest() for path in tmp_path.glob("*.docx")}
    scan_options = dict(initial_meeting_at=datetime(2026, 8, 11, 8),
                        final_meeting_at=datetime(2026, 8, 17, 10))
    episode = scan_patient_folder(tmp_path, **scan_options)
    discharge = scan_discharge_summary(tmp_path)
    assert discharge.icf_domains and discharge.discharge_scale_rows
    assert not any(issue.code == "final_mdrk_source_missing" for issue in discharge.issues)

    def export(document, name):
        path = tmp_path / name
        if document == "mdrk2":
            return write_mdrk_docx(scan_patient_folder(tmp_path, **scan_options), MdrkKind.FINAL, path)
        return write_discharge_summary_docx(scan_discharge_summary(tmp_path), path)

    second = "discharge" if first == "mdrk2" else "mdrk2"
    outputs = {first: export(first, first + ".docx")}
    outputs[second] = export(second, second + ".docx")
    export("mdrk2", "mdrk2-copy.docx")
    assert scan_discharge_summary(tmp_path) == discharge
    assert scan_patient_folder(tmp_path, **scan_options) == episode
    for document, path in outputs.items():
        repeated = export(document, "repeated-" + document + ".docx")
        assert Document(repeated).element.xml == Document(path).element.xml
    assert all(hashlib.sha256(path.read_bytes()).digest() == digest
               for path, digest in source_hashes.items())


@pytest.mark.parametrize("origin", ["mis", "final_clinician"])
def test_explicit_goal_result_does_not_require_mdrk2(tmp_path, origin):
    _write_document(tmp_path / "primary.docx", (*_primary_lines(),
        "Цель на этап медицинской реабилитации: ИСХОДНАЯ ЦЕЛЬ"))
    mis_path = tmp_path / "mis.docx"
    _write_document(mis_path, _discharge_lines())
    assert scan_discharge_summary(tmp_path).goal_result == ""
    goal = "Достигнута частично"
    result_lines = ("Состояние при выписке: улучшение",
                    "Цель, поставленная на этап медицинской реабилитации: " + goal,
                    "Трудоспособность: временно утрачена")
    if origin == "mis":
        source = mis_path
        _write_document(source, (*_discharge_lines(), *result_lines))
    else:
        source = tmp_path / "final.docx"
        _write_document(source, (
            "Заключительный осмотр невролога 16.08.2026 10:00",
            "ФИО пациента: Пациент Тестовый Пример", "Номер ИБ: СКП5906/26",
            *result_lines,
        ))
    draft = scan_discharge_summary(tmp_path)
    assert draft.goal_result == goal
    assert draft.field_sources["goal_result"] == source
    assert draft.final_mdrk_source is None
