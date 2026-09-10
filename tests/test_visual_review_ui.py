from copy import deepcopy
from datetime import datetime

from docx import Document

from test_document_panel_rescan import root
from mdrk_builder.domain import DischargeSummaryDraft, DischargeTeamFinding, SpecialistRole, Episode, MdrkKind
from mdrk_builder.infrastructure.draft_store import encode, decode
from mdrk_builder.infrastructure.discharge_summary_writer import write_discharge_summary_docx
from mdrk_builder.ui.discharge_summary_panel import DischargeSummaryPanel
from mdrk_builder.ui.formatted_text import FormattedText


def test_bold_selection_is_visual_persistent_and_rendered_without_markers(root, tmp_path):
    widget = FormattedText(root, undo=True)
    widget.insert("1.0", "Первая строка\nВторая строка")
    widget.tag_add("sel", "1.0", "2.6")
    widget.toggle_bold()
    value = widget.get("1.0", "end-1c")
    assert value == "**Первая строка**\n**Вторая** строка"
    widget.delete("1.0", "end")
    widget.insert("1.0", value)
    assert len(widget.tag_ranges("bold")) == 4
    draft = decode(encode(DischargeSummaryDraft(tmp_path, clinical_diagnosis=value)))
    result = Document(write_discharge_summary_docx(draft, tmp_path / "formatted.docx"))
    assert all("**" not in p.text for p in result.paragraphs)
    assert any(r.text == "Первая строка" and r.bold for p in result.paragraphs for r in p.runs)
    assert any(r.text == "Вторая" and r.bold for p in result.paragraphs for r in p.runs)


def test_discharge_specialists_keep_independent_edits_across_selection_and_rescan(root, tmp_path):
    panel = DischargeSummaryPanel(root, open_path=lambda _: None)
    draft = DischargeSummaryDraft(tmp_path, team_findings=(
        DischargeTeamFinding(SpecialistRole.NEUROLOGIST, "Врачебное заключение"),
        DischargeTeamFinding(SpecialistRole.PHYSICAL_THERAPIST, "Заключение ФР")))
    fresh = deepcopy(draft)
    panel.load(draft)
    editor = panel.specialists
    editor.text.delete("1.0", "end")
    editor.text.insert("1.0", "Ручная правка врача")
    editor.name.insert(0, "РУЧНОЕ ИМЯ")
    editor.list.selection_set("1")
    editor._select()
    assert editor.text.get("1.0", "end-1c") == "Заключение ФР"
    assert panel.draft.team_findings[0].conclusion == "Ручная правка врача"
    assert panel.draft.team_findings[0].specialist_name == "РУЧНОЕ ИМЯ"
    editor.text.insert("end-1c", " **с уточнением**")
    assert panel.merge_scan(fresh)
    assert panel.draft.team_findings[0].conclusion == "Ручная правка врача"
    assert panel.draft.team_findings[1].conclusion == "Заключение ФР **с уточнением**"
    editor._add_scale()
    editor._scale("0", "name", "Ручная шкала")
    editor._scale("0", "value", "8")
    assert panel.draft.team_findings[1].scales[0].value == "8"
    assert not panel.draft.team_findings[0].scales
    assert editor._sources.links("0") == editor._scale_links("0")
    assert any(label.startswith("Ручная правка") for label, _ in editor._scale_links("0"))


def test_mdrk_date_rebuild_uses_calendar_cutoff_and_preserves_late_text(root, tmp_path, monkeypatch):
    from mdrk_builder.ui.app import MdrkBuilderApp
    app = MdrkBuilderApp(root)
    episode = Episode(tmp_path, admission_datetime=datetime(2026, 8, 1),
        initial_meeting_at=datetime(2026, 8, 2, 8), final_meeting_at=datetime(2026, 8, 16, 8))
    app.episode = episode
    app._scan_baseline = deepcopy(episode)
    app._set_folder_field(str(tmp_path))
    app._populate_from_episode()
    captured, jobs = [], []
    def scan(folder, **kwargs):
        captured.append(kwargs)
        result = deepcopy(episode)
        result.initial_meeting_at = kwargs["initial_meeting_at"]
        result.final_meeting_at = kwargs["final_meeting_at"]
        return result
    monkeypatch.setattr("mdrk_builder.ui.document_dates.scan_patient_folder", scan)
    app._start_background_job = lambda operation, finish, **_: jobs.append((operation, finish))
    app._entry_variables["meeting"].set("03.08.2026 14:00")
    app._date_rebuilds.queue._pump()
    app._text_fields["clinical_diagnosis"].insert("1.0", "Врач дописал во время сканирования")
    app._dirty_section_fields[MdrkKind.INITIAL].add("clinical_diagnosis")
    operation, finish = jobs.pop()
    finish(operation(), None)
    assert captured[0]["initial_meeting_at"].date() == datetime(2026, 8, 3).date()
    assert captured[0]["initial_meeting_at"].hour == 23
    assert app.episode.meeting_at(MdrkKind.INITIAL) == datetime(2026, 8, 3, 14)
    assert app.episode.assessment_at(MdrkKind.INITIAL).day == 3
    assert app._text_fields["clinical_diagnosis"].get("1.0", "end-1c").startswith("Врач дописал")
    app._entry_variables["meeting"].set("03.08.2026 15:00")
    assert not app._date_rebuilds.queue.pending


def test_discharge_date_changes_during_scan_then_save_uses_latest_period(root, tmp_path, monkeypatch):
    from mdrk_builder.ui.app import MdrkBuilderApp
    app = MdrkBuilderApp(root)
    app._set_folder_field(str(tmp_path))
    app.document_var.set("discharge")
    start, old_end = datetime(2026, 8, 1), datetime(2026, 8, 9)
    panel = app.discharge_workspace
    panel.load(DischargeSummaryDraft(tmp_path, admission_datetime=start, discharge_datetime=old_end,
                                    projection_period=(start, old_end)))
    jobs, saved = [], []
    app._start_background_job = lambda operation, finish, **_: jobs.append((operation, finish))
    def scan(folder, **kwargs):
        end = kwargs["discharge_datetime_override"]
        return DischargeSummaryDraft(folder, admission_datetime=start, discharge_datetime=end,
            projection_period=(start, end), clinical_diagnosis="SOURCE DIAGNOSIS")
    monkeypatch.setattr("mdrk_builder.ui.document_dates.scan_discharge_summary", scan)
    app._save_auxiliary_document = lambda save: saved.append(deepcopy(panel.draft))
    panel._identity_vars["discharge"].set("11.08.2026")
    app._generate()
    app._date_rebuilds.queue._pump()
    panel._identity_vars["discharge"].set("12.08.2026")
    panel._widgets["clinical_diagnosis"].insert("1.0", "MANUAL DURING SCAN")
    operation, finish = jobs.pop()
    finish(operation(), None)
    assert not saved and panel.draft.discharge_datetime == old_end
    app._date_rebuilds.queue._pump()
    operation, finish = jobs.pop()
    finish(operation(), None)
    app._date_rebuilds.queue._pump()
    assert len(saved) == 1 and not saved[0].requires_period_rescan()
    assert saved[0].discharge_datetime.day == 12
    assert saved[0].clinical_diagnosis == "MANUAL DURING SCAN"


def test_quick_phrase_is_explicit_and_status_mode_survives_draft_and_rescan(root, tmp_path):
    from mdrk_builder.ui.quick_phrases import PHRASES, insert_phrase
    panel = DischargeSummaryPanel(root, open_path=lambda _: None)
    panel.load(DischargeSummaryDraft(tmp_path, neurological_status="NEURO", local_status="LOCAL"))
    assert panel._widgets["laboratory_results"].get("1.0", "end-1c") == ""
    insert_phrase(panel._widgets["laboratory_results"], PHRASES["laboratory_results"][0])
    panel.combine_statuses.set(True)
    panel._change_status_mode()
    assert panel.merge_scan(DischargeSummaryDraft(tmp_path, neurological_status="NEURO", local_status="LOCAL"))
    restored = decode(encode(panel.capture_state()))
    panel.clear()
    panel.restore_state(restored)
    assert panel.combine_statuses.get()
    assert panel.draft.neurological_status == "NEURO" and panel.draft.local_status == "LOCAL"
    assert panel.draft.laboratory_results == "Не проводились."
