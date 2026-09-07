from copy import deepcopy
from dataclasses import replace
from datetime import date, datetime
from pathlib import Path

import pytest

from mdrk_builder.application.diagnosis import diagnosis_parts, compose_diagnosis, LABELS
from mdrk_builder.application.editing import merge_rows, hospitalization_days
from mdrk_builder.application.conflicts import resolve_discharge_scale
from mdrk_builder.application.shared_edits import synchronize_discharge_point, synchronize_scale_rows, transfer_identity
from mdrk_builder.application.discharge_validation import current_discharge_issues
from mdrk_builder.application.discharge_identity import synchronize_header
from mdrk_builder.application.validation import generation_issues, current_issues
from mdrk_builder.domain import (DischargeScaleRow, DischargeTeamFinding, DischargeSummaryDraft,
    Episode, MdrkKind, Procedure, ReverseSheetDraft, ReverseSheetRow, SpecialistRole, SpecialistFinding, ScaleMeasurement)
from mdrk_builder.infrastructure.draft_store import load_draft, save_draft, decode
from mdrk_builder.application.scan_session import ScanSession, ScanCancelled
from mdrk_builder.application.source_scan import scan_source_documents
from mdrk_builder.ui.edit_history import EditHistory
from mdrk_builder.ui.reverse_sheet_dialog import incomplete_reverse_date_issues

ROLE = SpecialistRole.PHYSICAL_THERAPIST
START = datetime(2026, 8, 10)
END = datetime(2026, 8, 18)


def test_diagnosis_three_parts_fallback_and_conflict():
    primary = Path('primary.docx')
    mis = Path('mis.docx')
    value, sources, issues = compose_diagnosis([
        (primary, 'Основное заболевание\nОсновное А\nДополнительные сведения о заболевании\nСведения А'),
        (mis, 'Основное заболевание: Основное Б\nСопутствующие заболевания: Сопутствующее Б')])
    assert diagnosis_parts(value) == {'main': 'Основное А', 'additional': 'Сопутствующее Б', 'factors': 'Сведения А'}
    assert sources['clinical_diagnosis.main'] == primary
    assert sources['clinical_diagnosis.additional'] == mis
    assert len(issues) == 1


def test_three_way_merge_retains_new_rows_and_unedited_source_fields():
    base = [Procedure('А', 'ФТ', 3, source=Path('a.docx'))]
    edited = deepcopy(base)
    edited[0].actual_count = 5
    edited[0].manual_fields.add('actual_count')
    incoming = [replace(base[0], actual_count=4, frequency='ежедневно'), Procedure('Б','ФТ',2)]
    result, conflicts = merge_rows(base, edited, incoming)
    assert [r.name for r in result] == ['А', 'Б']
    assert result[0].actual_count == 5
    assert result[0].frequency == 'ежедневно'
    assert conflicts
    assert not incoming[0].manual_fields
    assert not merge_rows(base, [], base)[0]


def test_resolving_current_scale_does_not_overwrite_admission():
    draft = DischargeSummaryDraft(Path('.'))
    initial = DischargeScaleRow(ROLE, 'Берг', '20', Path('a.docx'), current_at=START)
    final = replace(initial, value='30', source=Path('b.docx'), current_at=END)
    pair = replace(final, initial_value='20', initial_source=initial.source, initial_at=START)
    draft.admission_scale_rows = (initial,)
    draft.discharge_scale_rows = (final,)
    draft.team_findings = (DischargeTeamFinding(ROLE, '', scales=(pair,)),)
    resolve_discharge_scale(draft, f'scale:{ROLE.value}|{END.isoformat()}|Берг', '35', Path('c.docx'))
    assert draft.admission_scale_rows[0].value == '20'
    assert draft.discharge_scale_rows[0].value == '35'
    assert draft.team_findings[0].scales[0].initial_value == '20'
    assert draft.team_findings[0].scales[0].value == '35'


def test_general_scale_edits_and_deletion_update_specialist_table():
    draft = DischargeSummaryDraft(Path('.'))
    initial = DischargeScaleRow(ROLE, 'Берг', '20', current_at=START)
    final = replace(initial, value='35', current_at=END)
    synchronize_discharge_point(draft, 'admission_scale_rows', None, initial)
    synchronize_discharge_point(draft, 'discharge_scale_rows', None, final)
    row = draft.team_findings[0].scales[0]
    assert (row.initial_value, row.value) == ('20', '35')
    synchronize_scale_rows(draft, row)
    assert draft.admission_scale_rows[0].current_at == START
    assert draft.discharge_scale_rows[0].current_at == END
    synchronize_discharge_point(draft, 'discharge_scale_rows', final, None)
    assert draft.team_findings[0].scales[0].value == ''
    assert draft.team_findings[0].scales[0].initial_value == '20'


def test_header_identity_age_and_days():
    draft = DischargeSummaryDraft(Path('.'), admission_datetime=START, discharge_datetime=END,
        header_text='ФИО: старое\nДата рождения: старое\nПол: старое\nНомер медицинской карты: старое\nДругая строка')
    draft.identity.full_name = 'Тестовый Пациент'
    draft.identity.birth_date = date(2005,8,11)
    draft.identity.medical_record_number = '123'
    synchronize_header(draft)
    assert 'Тестовый Пациент' in draft.header_text
    assert '21 год' in draft.header_text
    assert 'Другая строка' in draft.header_text
    assert 'старое' not in draft.header_text
    assert hospitalization_days(START, END) == 8
    assert hospitalization_days(START, START) == 1
    assert hospitalization_days(END, START) is None
    assert '8' in draft.header_text


def test_validation_refreshes_diagnosis_and_ranges_and_counts():
    draft = DischargeSummaryDraft(Path('.'), discharge_datetime=END, header_text='Шапка')
    assert any(i.field == 'clinical_diagnosis' for i in current_discharge_issues(draft))
    draft.clinical_diagnosis = '\n'.join(v + ':' for v in LABELS.values())
    assert any(i.field == 'clinical_diagnosis' for i in current_discharge_issues(draft))
    draft.clinical_diagnosis = 'Основное заболевание: диагноз'
    draft.admission_scale_rows = (DischargeScaleRow(ROLE, 'Берг','999'),)
    draft.completed_procedures = (Procedure('ЛФК','ФТ',3,performed_dates=(date(2026,8,10),)),)
    codes = {i.code for i in current_discharge_issues(draft)}
    assert 'discharge_current_required' not in codes
    assert {'scale_value_out_of_range','procedure_dates_count_mismatch'} <= codes
    draft.issues = current_discharge_issues(draft)
    draft.admission_scale_rows = (replace(draft.admission_scale_rows[0],value='30'),)
    assert 'scale_value_out_of_range' not in {i.code for i in current_discharge_issues(draft)}


def test_ranges_checked_at_both_mdrk_validation_entrypoints():
    episode = Episode(Path('.'))
    episode.findings = [SpecialistFinding(ROLE, scales=[ScaleMeasurement('Берг', '999', START, ROLE)])]
    for validate in (generation_issues, current_issues):
        assert 'scale_value_out_of_range' in {i.code for i in validate(episode, MdrkKind.FINAL)}


def test_reverse_chronology_and_shared_identity():
    row = ReverseSheetRow('ЛФК', date(2026,8,11), START)
    assert {'reverse_date_order','reverse_outside_period'} <= {i.code for i in incomplete_reverse_date_issues([row], END, END)}
    discharge = DischargeSummaryDraft(Path('.'), discharge_datetime=END)
    discharge.identity.full_name = 'Новое имя'
    reverse = ReverseSheetDraft(Path('.'))
    transfer_identity(discharge, reverse, {'full_name','discharge'})
    assert reverse.discharge_datetime == END
    assert reverse.identity.full_name == 'Новое имя'


def test_rich_draft_roundtrip_is_data_only(tmp_path):
    draft = DischargeSummaryDraft(tmp_path, discharge_datetime=END)
    draft.team_findings = (DischargeTeamFinding(ROLE,'Заключение',scales=(DischargeScaleRow(ROLE,'Берг','30',manual_fields={'value'}),)),)
    state = {'draft':draft, 'kind':MdrkKind.FINAL, 'dirty':{'name'}, 'keys':{MdrkKind.INITIAL:'x'}}
    path = tmp_path / '.mdrk draft.json'
    save_draft(path,state)
    assert load_draft(path) == state
    with pytest.raises(ValueError):
        decode({'type':'exec','fields':{}})


def test_session_reuses_read_and_invalidates_changed_content(tmp_path, monkeypatch):
    from docx import Document
    from mdrk_builder.application import source_scan
    path = tmp_path / 'source.docx'
    doc = Document();doc.add_paragraph('Первичный осмотр невролога');doc.save(path)
    original = source_scan.read_docx
    reads = []
    def counted(*args, **kwargs):
        reads.append(args[0]);return original(*args, **kwargs)
    monkeypatch.setattr(source_scan,'read_docx',counted)
    session = ScanSession();session.begin(tmp_path)
    scan_source_documents(tmp_path,session=session)
    scan_source_documents(tmp_path,session=session)
    assert len(reads)==1
    doc.add_paragraph('Изменение');doc.save(path)
    scan_source_documents(tmp_path,session=session)
    assert len(reads)==2
    session.cancelled.set()
    with pytest.raises(ScanCancelled):scan_source_documents(tmp_path,session=session)
    session.begin(tmp_path/'other')
    assert not session.cache


def test_table_undo_redo_and_history_reset():
    state = [1]
    history = EditHistory(lambda:state,lambda v:state.__setitem__(slice(None),v))
    history.wrap(lambda:state.append(2))()
    history.undo();assert state==[1]
    history.redo();assert state==[1,2]
    history.clear();history.undo();assert state==[1,2]


def test_merge_keeps_two_dates_in_one_source():
    rows = [ScaleMeasurement('Берг','20',START,ROLE,Path('one.docx')),
            ScaleMeasurement('Берг','30',END,ROLE,Path('one.docx'))]
    edited = deepcopy(rows);edited[0].value='21'
    merged,_ = merge_rows(rows,edited,rows)
    assert [row.value for row in merged]==['21','30']


def test_shared_scale_changes_from_one_file_are_dated_and_marked():
    from mdrk_builder.application.shared_edits import transfer_episode_edits, transfer_discharge_edits, remove_scale_rows
    source=Path('both.docx')
    episode=Episode(Path('.'), findings=[SpecialistFinding(ROLE,scales=[ScaleMeasurement('Берг','20',START,ROLE,source), ScaleMeasurement('Берг','30',END,ROLE,source)])])
    baseline=deepcopy(episode)
    draft=DischargeSummaryDraft(Path('.'), admission_scale_rows=(DischargeScaleRow(ROLE,'Берг','20',source,current_at=START),),discharge_scale_rows=(DischargeScaleRow(ROLE,'Берг','30',source,current_at=END),))
    episode.findings[0].scales[0].value='21';episode.findings[0].scales[0].manual_fields.add('value')
    transfer_episode_edits(episode,baseline,draft)
    assert [draft.admission_scale_rows[0].value,draft.discharge_scale_rows[0].value]==['21','30']
    draft.discharge_scale_rows=(replace(draft.discharge_scale_rows[0],value='35',manual_fields={'value'}),)
    assert 'findings' in transfer_discharge_edits(draft,episode)
    assert [s.value for s in episode.findings[0].scales]==['21','35']
    draft.discharge_scale_rows += (DischargeScaleRow(ROLE,'Берг','40',source,current_at=datetime(2026,8,19)),)
    pair=replace(draft.discharge_scale_rows[0],initial_value='21',initial_at=START,initial_source=source)
    remove_scale_rows(draft,pair)
    assert [s.value for s in draft.discharge_scale_rows]==['40']


def test_manual_scale_addition_and_date_change_transfer_to_mdrk():
    from mdrk_builder.application.shared_edits import transfer_discharge_edits
    draft=DischargeSummaryDraft(Path('.'), admission_scale_rows=(DischargeScaleRow(ROLE,'Берг','20',current_at=START,manual_fields={'value'}),))
    episode=Episode(Path('.'))
    transfer_discharge_edits(draft,episode)
    assert len(episode.findings[0].scales)==1
    baseline=deepcopy(draft)
    draft.admission_scale_rows=(replace(draft.admission_scale_rows[0],current_at=END),)
    transfer_discharge_edits(draft,episode,baseline=baseline)
    assert [r.measured_at for r in episode.findings[0].scales]==[END]
    baseline=deepcopy(draft)
    draft.admission_scale_rows=()
    transfer_discharge_edits(draft,episode,baseline=baseline)
    assert not episode.findings[0].scales
