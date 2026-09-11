from dataclasses import replace
from datetime import datetime

from docx import Document

from mdrk_builder.application.discharge_current_fields import select_current_fields
from mdrk_builder.application.reverse_sheet import scan_reverse_sheet
from mdrk_builder.domain import DischargeScaleRow, DischargeSummaryDraft, DischargeTeamFinding, SpecialistRole
from mdrk_builder.infrastructure.discharge_summary_writer import write_discharge_summary_docx
from mdrk_builder.infrastructure.classifier import DocumentClassification
from test_discharge_summary import _scanned_source, _discharge_lines, _primary_lines, _write_document


def test_specialist_tables_have_header_dates_and_no_duplicate_physician_scores(tmp_path):
    doctor, motor = SpecialistRole.NEUROLOGIST, SpecialistRole.PHYSICAL_THERAPIST
    start, middle, end = datetime(2026, 8, 1), datetime(2026, 8, 9), datetime(2026, 8, 11)
    scale = DischargeScaleRow(doctor, "Врачебная шкала", "4", initial_value="3", initial_at=start, current_at=middle)
    motor_scale = replace(scale, role=motor, name="Двигательная шкала")
    draft = DischargeSummaryDraft(tmp_path, admission_datetime=start, discharge_datetime=end,
        admission_scale_rows=(replace(scale, value="3"),), discharge_scale_rows=(scale,),
        team_findings=(DischargeTeamFinding(doctor, "Итог врача", scales=(scale,),
                        specialist_title="Врач физической и реабилитационной медицины"),
                       DischargeTeamFinding(motor, "Итог ФР", occurred_at=end, scales=(motor_scale,))))
    document = Document(write_discharge_summary_docx(draft, tmp_path / "scales.docx"))
    rows = [row for table in document.tables for row in table.rows[1:]]
    assert sum(row.cells[0].text == "Врачебная шкала" for row in rows) == 2
    motor_table = next(t for t in document.tables if any(r.cells[0].text == "Двигательная шкала" for r in t.rows))
    assert "11.08.2026" in motor_table.rows[0].cells[2].text
    assert [c.text for c in motor_table.rows[1].cells] == ["Двигательная шкала", "3", "4"]
    assert not any(c.text == "Специалист" for t in document.tables for c in t.rows[0].cells)
    assert any(p.text.startswith("Результат осмотра врача физической и реабилитационной медицины (") for p in document.paragraphs)
    assert any(p.text.startswith("Результат осмотра специалиста по физической реабилитации (") for p in document.paragraphs)


def test_combined_status_keeps_both_texts_and_removes_only_unwanted_header_lines(tmp_path):
    draft = DischargeSummaryDraft(tmp_path, header_text="ФИО: ТЕСТ\nМодель: служебная\nНастоящая госпитализация: текст\nНомер ИБ: 7",
        neurological_status="НЕВРОЛОГИЧЕСКИЙ", local_status="ЛОКАЛЬНЫЙ", combine_admission_statuses=True,
        clinical_diagnosis="Основное заболевание:\nДИАГНОЗ\nСопутствующие заболевания: СОПУТСТВУЮЩЕЕ")
    document = Document(write_discharge_summary_docx(draft, tmp_path / "status.docx"))
    paragraphs = [p.text for p in document.paragraphs]
    assert not any(p.startswith(("Модель", "Настоящая госпитализация", "Локальный статус:")) for p in paragraphs)
    assert "Неврологический и локальный статус: НЕВРОЛОГИЧЕСКИЙ\nЛОКАЛЬНЫЙ" in paragraphs
    diagnosis = next(p for p in document.paragraphs if p.text.startswith("Основное заболевание:"))
    assert diagnosis.text == "Основное заболевание: ДИАГНОЗ" and diagnosis.runs[0].bold


def test_negative_mis_field_retains_contradicting_source_for_physician_choice():
    primary = _scanned_source("primary.docx", *_primary_lines(), "Дата осмотра: 10.08.2026 11:00", document_type="initial")
    primary = replace(primary, classification=DocumentClassification(SpecialistRole.NEUROLOGIST, "initial"))
    mis = _scanned_source("mis.docx", *_discharge_lines(),
        "Факторы, ограничивающие проведение реабилитационных мероприятий: нет", "| |", document_type="discharge_summary")
    values, sources, choices, issues = select_current_fields((primary, mis), datetime(2026, 8, 10), datetime(2026, 8, 17, 23), mis.document.source_path)
    assert values["limitations"] == "нет"
    assert {v for v, _ in choices["limitations"]} == {"нет", "PRIMARY LIMITS"}
    assert any(issue.field == "limitations" for issue in issues)
    assert sources["limitations"] == mis.document.source_path


def test_future_discharge_date_does_not_redate_mis_or_exclude_newer_specialist_fact():
    mis = _scanned_source("mis.docx", *_discharge_lines(),
        "Факторы, ограничивающие проведение реабилитационных мероприятий: нет", document_type="discharge_summary")
    motor = _scanned_source("motor.docx", "Повторный осмотр специалиста по физической реабилитации",
        "Дата осмотра: 18.08.2026 11:00", "Факторы, ограничивающие проведение реабилитационных мероприятий: ОГРАНИЧЕНИЕ")
    motor = replace(motor, classification=DocumentClassification(SpecialistRole.PHYSICAL_THERAPIST, "follow_up"))
    values, sources, _, _ = select_current_fields((mis, motor), datetime(2026, 8, 10), datetime(2026, 8, 19), mis.document.source_path)
    assert values["limitations"] == "ОГРАНИЧЕНИЕ"
    assert sources["limitations"] == motor.document.source_path


def test_reverse_sheet_respects_both_calendar_boundaries(tmp_path):
    _write_document(tmp_path / "primary.docx", _primary_lines())
    for day in (11, 15, 18):
        _write_document(tmp_path / f"speech-{day}.docx", ("Повторный осмотр медицинского логопеда",
            f"Дата осмотра: {day}.08.2026 17:00", "Медицинский логопед: ТЕСТ А.А."))
    draft = scan_reverse_sheet(tmp_path, admission_datetime_override=datetime(2026, 8, 12), discharge_datetime_override=datetime(2026, 8, 15))
    assert [row.performed_at.day for row in draft.rows if row.performed_at] == [15]


def test_numeric_hospitalization_period_keeps_source_discharge_time():
    from mdrk_builder.application.discharge_extractors import extract_summary_discharge_datetime
    source = _scanned_source("mis.docx", "Период нахождения в стационаре: с 01.08.2026 10:00 по 17.08.2026 12:00")
    assert extract_summary_discharge_datetime(source.document) == datetime(2026, 8, 17, 12)


def test_last_available_score_retains_its_source_date_and_independent_edit(tmp_path):
    from mdrk_builder.domain import Episode, MdrkKind, ScaleMeasurement, SpecialistFinding
    from mdrk_builder.application.snapshot import build_snapshot
    from mdrk_builder.application.discharge_summary import _project_scale_rows, _project_team_findings
    from mdrk_builder.ui.discharge_table_fields import edit_field
    at = datetime(2026, 8, 1)
    measurement = ScaleMeasurement("Berg", "12", at, SpecialistRole.PHYSICAL_THERAPIST, tmp_path / "initial.docx")
    episode = Episode(tmp_path, initial_meeting_at=at, final_meeting_at=datetime(2026, 8, 16),
        findings=[SpecialistFinding(measurement.specialist, source_datetime=at, source=measurement.source, scales=[measurement])])
    snapshot = build_snapshot(episode, MdrkKind.FINAL)
    admission, discharge = _project_scale_rows(snapshot, final_mdrk_source=None)
    assert admission[0].value == discharge[0].value == "12"
    assert discharge[0].current_at == at and discharge[0].source == measurement.source
    pair = _project_team_findings(snapshot)[0].scales[0]
    edited = edit_field(pair, "value", "14")
    assert edited.initial_value == "12" and edited.value == "14"
    assert measurement.value == "12" and measurement.measured_at == at
