"""Document formatting regressions using synthetic episode data only."""
from datetime import datetime

import pytest
from docx import Document

from mdrk_builder.application.discharge_summary import _project_team_findings
from mdrk_builder.application.snapshot import build_snapshot
from mdrk_builder.domain import (
    DischargeSummaryDraft, DischargeTeamFinding, IcfDomain, IcfQualifier, MdrkKind,
    SourceDocument, SpecialistFinding, SpecialistRole,
)
from mdrk_builder.infrastructure.discharge_summary_writer import write_discharge_summary_docx
from mdrk_builder.infrastructure.docx_writer import write_mdrk_docx
from test_docx_writer import _cell_fill, _find_domain_row, _find_table, _representative_episode


@pytest.mark.parametrize("output_kind", ["initial", "final", "discharge"])
def test_icf_shading_uses_the_assessment_for_the_document(tmp_path, output_kind):
    episode = _representative_episode(tmp_path)
    cases = [
        ("b730", IcfQualifier(3), IcfQualifier(2), {6, 7, 8}),
        ("d450", IcfQualifier(1), IcfQualifier(3), {6, 7, 8, 9}),
        ("b130", IcfQualifier(3), IcfQualifier(0), {6}),
        ("e310", IcfQualifier(4, True), IcfQualifier(2, True), {4, 5, 6}),
        ("d640", IcfQualifier(2), None, {6, 7, 8}),
        ("b140", None, IcfQualifier(1), {6, 7}),
    ]
    episode.icf_domains = [
        IcfDomain(code, "ТЕСТОВЫЙ ДОМЕН", SpecialistRole.OTHER, initial=initial, final=final)
        for code, initial, final, _ in cases
    ]
    output = tmp_path / (output_kind + ".docx")
    if output_kind == "discharge":
        draft = DischargeSummaryDraft(tmp_path, icf_domains=tuple(episode.icf_domains))
        write_discharge_summary_docx(draft, output)
    else:
        write_mdrk_docx(episode, MdrkKind(output_kind), output, ignore_issues=True)
    table = _find_table(Document(output), "МКФ категориальный профиль")
    initial_columns = [{6, 7, 8, 9}, {6, 7}, {6, 7, 8, 9}, {2, 3, 4, 5, 6}, {6, 7, 8}, set()]
    for (code, initial, final, expected_final), expected_initial in zip(cases, initial_columns):
        if output_kind == "initial" and initial is None:
            assert not any(row.cells[0].text == code for row in table.rows)
            continue
        row = _find_domain_row(table, code)
        expected = expected_initial if output_kind == "initial" else expected_final
        assert {i for i in range(2, 11) if _cell_fill(row.cells[i]) == "BFBFBF"} == expected
        assert row.cells[11].text == (initial.display() if initial else "")
        if output_kind == "discharge":
            assert row.cells[12].text == (final.display() if final else "")


@pytest.mark.parametrize("kind", [MdrkKind.INITIAL, MdrkKind.FINAL])
def test_mdrk_scale_headers_contain_only_dates(tmp_path, kind):
    episode = _representative_episode(tmp_path)
    document = Document(write_mdrk_docx(episode, kind, tmp_path / "mdrk.docx"))
    table = next(t for t in document.tables if any(r.cells[0].text == "Шкала Тинетти" for r in t.rows))
    expected = ["Шкала/опросник", "05.06.2026"]
    if kind is MdrkKind.FINAL:
        expected.append("19.06.2026")
    assert [cell.text for cell in table.rows[0].cells] == expected


@pytest.mark.parametrize("output_kind", ["initial", "final", "discharge"])
def test_result_headings_include_the_selected_exam_name_title_and_timestamp(tmp_path, output_kind):
    episode = _representative_episode(tmp_path)
    role = SpecialistRole.LOGOPEDIST
    initial_at, final_at = datetime(2026, 6, 5, 11, 20), datetime(2026, 6, 18, 14, 35)
    episode.sources += [
        SourceDocument(tmp_path / "speech-initial.docx", role, initial_at, specialist_name="СОТРУДНИК_ПЕРВЫЙ И.И."),
        SourceDocument(tmp_path / "speech-final.docx", role, final_at, specialist_name="СОТРУДНИК_ПОВТОРНЫЙ П.П."),
    ]
    episode.findings = [
        SpecialistFinding(role, "ПЕРВИЧНЫЙ ОСМОТР", initial_at, episode.sources[-2].path,
                          specialist_title="Медицинский логопед"),
        SpecialistFinding(role, "ПОВТОРНЫЙ ОСМОТР", final_at, episode.sources[-1].path,
                          specialist_title="Медицинский логопед"),
    ]
    output = tmp_path / (output_kind + ".docx")
    if output_kind == "discharge":
        snapshot = build_snapshot(episode, MdrkKind.FINAL)
        draft = DischargeSummaryDraft(tmp_path, team_findings=_project_team_findings(snapshot, episode.sources))
        write_discharge_summary_docx(draft, output)
    else:
        write_mdrk_docx(episode, MdrkKind(output_kind), output)
    expected = (
        "Результат осмотра медицинского логопеда СОТРУДНИК_ПЕРВЫЙ И.И. (05.06.2026, 11:20)"
        if output_kind == "initial" else
        "Результат осмотра медицинского логопеда СОТРУДНИК_ПОВТОРНЫЙ П.П. (18.06.2026, 14:35)"
    )
    assert expected in [p.text for p in Document(output).paragraphs]


@pytest.mark.parametrize("output_kind", ["initial", "final", "discharge"])
@pytest.mark.parametrize("role", [SpecialistRole.NEUROLOGIST, SpecialistRole.FRM])
def test_physician_exam_is_always_presented_as_frm(tmp_path, output_kind, role):
    episode = _representative_episode(tmp_path)
    at = episode.initial_meeting_at
    source = tmp_path / "physician.docx"
    name = "СОТРУДНИК И.И."
    episode.sources = [SourceDocument(source, role, at, specialist_name=name)]
    episode.findings = [SpecialistFinding(role, "КОНТРОЛЬНЫЙ ОСМОТР", at, source,
                                          specialist_title="Врач-невролог")]
    output = tmp_path / (output_kind + ".docx")
    if output_kind == "discharge":
        draft = DischargeSummaryDraft(tmp_path, team_findings=(
            DischargeTeamFinding(role, "КОНТРОЛЬНЫЙ ОСМОТР", occurred_at=at,
                                 specialist_name=name, specialist_title="Врач-невролог"),))
        write_discharge_summary_docx(draft, output)
    else:
        write_mdrk_docx(episode, MdrkKind(output_kind), output)
    headings = [p.text for p in Document(output).paragraphs if p.text.startswith("Результат осмотра")]
    assert any(heading.startswith(f"Результат осмотра врача физической и реабилитационной медицины {name} (")
               for heading in headings)
    assert not any("невролог" in heading.casefold() for heading in headings)
    assert episode.findings[0].specialist_title == "Врач-невролог"
