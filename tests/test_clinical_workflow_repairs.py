from datetime import datetime
from pathlib import Path

from docx import Document
import pytest

from mdrk_builder.application.discharge_examinations import ExaminationSection, partition_examinations
from mdrk_builder.application.discharge_summary import scan_discharge_summary
from mdrk_builder.application.extractors import extract_clinical_sections, extract_conclusion
from mdrk_builder.application.scanner import scan_patient_folder
from mdrk_builder.application.snapshot import build_snapshot
from mdrk_builder.domain import MdrkKind, SpecialistRole
from mdrk_builder.infrastructure.ooxml_reader import read_docx
from test_discharge_summary import _primary_lines, _discharge_lines, _write_document

NAME = "АЛЬФА БЕТА ГАММА"


def test_previous_hospitalization_cannot_replace_current_without_mis(tmp_path):
    primary = tmp_path / "primary.docx"
    old = tmp_path / "previous.docx"
    _write_document(primary, (*_primary_lines(full_name=NAME), "Дата осмотра: 10.08.2026 11:00"))
    document = Document(primary)
    table = document.add_table(rows=2, cols=15)
    table.cell(0, 0).text = "МКФ"
    table.cell(1, 0).text = "s110"
    table.cell(1, 1).text = "Структура головного мозга"
    table.cell(1, 11).text = "2"
    document.save(primary)
    _write_document(old, (*_discharge_lines("СКП1000/26", full_name=NAME),
        "Состояние при выписке: OLD STATE", "Рекомендации: OLD ADVICE", "Лечащий врач: OLD SIGNATURE"))
    draft = scan_discharge_summary(tmp_path)
    assert draft.discharge_source is None
    assert draft.primary_neurologist_source == primary
    assert draft.identity.full_name == NAME
    assert draft.identity.medical_record_number == "СКП5906/26"
    assert "PRIMARY DIAGNOSIS" in draft.clinical_diagnosis
    assert draft.icf_domains
    assert all("OLD" not in getattr(draft, field) for field in
               ("header_text", "discharge_condition", "recommendations", "signatures"))
    assert old not in draft.field_sources.values()
    assert draft.rehabilitation_potential == "средний"
    assert draft.goal_result == "достигнут в полном объёме"


def test_confirmed_mis_fills_missing_demographics(tmp_path):
    _write_document(tmp_path / "primary.docx", _primary_lines(full_name=NAME))
    mis = tmp_path / "mis.docx"
    _write_document(mis, (*_discharge_lines(full_name=NAME), "Пол: мужской"))
    draft = scan_discharge_summary(tmp_path)
    assert draft.identity.sex == "мужской"
    assert draft.field_sources["identity.sex"] == mis


def test_latest_diary_supplies_discharge_examination_without_mis(tmp_path):
    _write_document(tmp_path / "primary.docx", _primary_lines(full_name=NAME))
    diary = tmp_path / "diary.docx"
    _write_document(diary, (
        "Повторный осмотр невролога 16.08.2026 10:00", f"ФИО пациента: {NAME}",
        "Номер ИБ: СКП5906/26", "Физикальное обследование: CURRENT PHYSICAL",
        "Неврологический статус: CURRENT NEUROLOGICAL", "Реабилитационный потенциал: высокий",
    ))
    draft = scan_discharge_summary(tmp_path)
    assert draft.discharge_condition == "CURRENT PHYSICAL"
    assert draft.discharge_neurological_status == "CURRENT NEUROLOGICAL"
    assert draft.field_sources["discharge_condition"] == diary
    assert draft.rehabilitation_potential == "высокий"


def test_initial_diagnosis_tasks_and_defaults(tmp_path):
    primary = tmp_path / "primary.docx"
    _write_document(primary, (*_primary_lines(full_name=NAME)[:6],
        "Диагноз:", "клинический:", "Основное заболевание: MAIN",
        "Сопутствующее заболевание: ACCOMPANYING", "Дополнительные сведения: EXTRA", "ШРМ: 3",
        "Реабилитационный диагноз:", "МКФ",))
    episode = scan_patient_folder(tmp_path)
    snapshot = build_snapshot(episode, MdrkKind.INITIAL)
    assert all(text in snapshot.sections.clinical_diagnosis for text in ("MAIN", "ACCOMPANYING", "EXTRA", "ШРМ: 3"))
    assert episode.course_duration(MdrkKind.INITIAL) == 16
    assert any(f.conclusion == "показано проведение курса реабилитационного лечения" for f in episode.findings)


def test_neurologist_default_does_not_replace_explicit_conclusion_with_empty_diary(tmp_path):
    _write_document(tmp_path / 'primary.docx', (*_primary_lines(full_name=NAME),
        'Дата осмотра: 10.08.2026 11:00', 'Заключение: EXPLICIT'))
    _write_document(tmp_path / 'diary.docx', (
        'Повторный осмотр невролога 16.08.2026 10:00', f'ФИО пациента: {NAME}',
        'Номер ИБ: СКП5906/26', 'ШРМ: 3'))
    snapshot = build_snapshot(scan_patient_folder(tmp_path), MdrkKind.FINAL)
    neurologist = next(row for row in snapshot.findings if row.role is SpecialistRole.NEUROLOGIST)
    assert neurologist.conclusion == 'EXPLICIT'


def test_shrm_outside_diagnosis_is_included_in_initial_diagnosis(tmp_path):
    path = tmp_path / 'primary.docx'
    _write_document(path, ('Первичный осмотр невролога', 'Диагноз: MAIN',
                           'Реабилитационный диагноз: END', 'ШРМ: 3'))
    assert extract_clinical_sections(read_docx(path))['clinical_diagnosis'] == 'MAIN\nШРМ: 3'


@pytest.mark.parametrize("role", [SpecialistRole.LOGOPEDIST, SpecialistRole.PATHOPSYCHOLOGIST])
def test_numbered_rehabilitation_tasks_exclude_short_goal(tmp_path, role):
    path = tmp_path / "specialist.docx"
    _write_document(path, (role.display_name, "11. Реабилитационные задачи на этап МР:",
        "1. TASK ONE", "Краткосрочная цель: EXCLUDE GOAL", "2. TASK TWO", "Рекомендации: ADVICE"))
    tasks = extract_clinical_sections(read_docx(path))["tasks"]
    assert "TASK ONE" in tasks and "TASK TWO" in tasks
    assert "EXCLUDE GOAL" not in tasks and "ADVICE" not in tasks


def test_neuropsych_full_rationale_and_procedure_tail(tmp_path):
    path = tmp_path / "psychologist.docx"
    rationale = "Исследование анамнеза, результатов обследования не обнаруживает оснований для продолжения работы с медицинским психологом."
    _write_document(path, ("Нейропсихологический статус: STATUS", rationale, "Рекомендовано:",
        "A13.29.001 ПРОЦЕДУРА", "(ПОЛНОЕ ПРОДОЛЖЕНИЕ)", "Медицинский психолог: SIGNATURE"))
    value = extract_conclusion(read_docx(path), SpecialistRole.NEUROPSYCHOLOGIST)
    assert rationale in value
    assert "A13.29.001 ПРОЦЕДУРА\n(ПОЛНОЕ ПРОДОЛЖЕНИЕ)" in value
    assert "SIGNATURE" not in value


def test_speech_recommendations_end_at_individual_program_label(tmp_path):
    path = tmp_path / "speech.docx"
    _write_document(path, ("Логопедический статус: STATUS", "На основании данных рекомендовано:",
        "1. ADVICE ONE", "2. ADVICE TWO", "3. Индивидуальная программа: преодоление DETAIL",
        "Восстановление MORE DETAIL", "Медицинский логопед: SIGNATURE"))
    value = extract_conclusion(read_docx(path), SpecialistRole.LOGOPEDIST)
    assert "ADVICE ONE" in value and "ADVICE TWO" in value
    assert "3. Индивидуальная программа." in value
    assert "DETAIL" not in value and "SIGNATURE" not in value


def test_examinations_are_partitioned_by_performance_date_and_keep_sources():
    fields = {"provided_documents": "ОАК от 11.08.2026 CURRENT LAB", "laboratory_results": "ОАК от 09.08.2026 BEFORE",
              "instrumental_results": "ЭКГ от 12.08.2026 CURRENT ECG\nЭКГ от 20.08.2026 AFTER"}
    sources = {name: Path(name + ".docx") for name in fields}
    sections = [ExaminationSection(name, text, sources[name]) for name, text in fields.items()]
    values, origins, issues = partition_examinations(sections, datetime(2026, 8, 10), datetime(2026, 8, 17))
    assert "BEFORE" in values["provided_documents"]
    assert "CURRENT LAB" in values["laboratory_results"]
    assert "CURRENT ECG" in values["instrumental_results"] and "AFTER" not in values["instrumental_results"]
    assert origins["laboratory_results"] == sources["provided_documents"]
    assert any(i.code == "examination_after_discharge" for i in issues)


@pytest.mark.parametrize('role,heading', [
    (SpecialistRole.NEUROPSYCHOLOGIST, 'Нейропсихологический статус'),
    (SpecialistRole.LOGOPEDIST, 'Логопедический статус'),
])
def test_recommendations_belong_to_selected_examination(tmp_path, role, heading):
    path = tmp_path / 'specialist.docx'
    _write_document(path, (f'{heading}: STATUS_A', 'Рекомендовано: ADVICE_A', 'Подпись: A',
                           f'{heading}: STATUS_B', 'Рекомендовано: ADVICE_B', 'Подпись: B'))
    value = extract_conclusion(read_docx(path), role)
    assert 'STATUS_B' in value and 'ADVICE_B' in value
    assert 'STATUS_A' not in value and 'ADVICE_A' not in value


def test_direct_speech_recommendations_do_not_keep_program_details_in_status(tmp_path):
    path = tmp_path / 'speech.docx'
    _write_document(path, ('Логопедический статус: STATUS', 'Рекомендовано:',
                           '1. ADVICE', '2. Индивидуальная программа: восстановление DETAIL',
                           'Преодоление MORE DETAIL', 'Медицинский логопед: SIGNATURE'))
    value = extract_conclusion(read_docx(path), SpecialistRole.LOGOPEDIST)
    assert value.count('Индивидуальная программа.') == 1
    assert 'DETAIL' not in value and 'ADVICE' in value


def test_earlier_speech_dynamics_do_not_override_latest_status(tmp_path):
    path = tmp_path / 'speech.docx'
    _write_document(path, ('Динамика: OLD', 'Логопедический статус: OLD',
        'Рекомендовано: OLD', 'Подпись: A', 'Логопедический статус: CURRENT',
        'Рекомендовано: CURRENT', 'Подпись: B'))
    value = extract_conclusion(read_docx(path), SpecialistRole.LOGOPEDIST)
    assert 'OLD' not in value and value.count('CURRENT') == 2


def test_examination_title_date_and_result_move_together():
    before = 'Общий анализ крови\n09.08.2026\nГемоглобин 123'
    current = 'ЭКГ\n12.08.2026\nСинусовый ритм'
    sections = [ExaminationSection('laboratory_results', before, Path('primary.docx')),
                ExaminationSection('provided_documents', current, Path('primary.docx'))]
    values, origins, issues = partition_examinations(sections, datetime(2026, 8, 10), datetime(2026, 8, 17))
    assert values['provided_documents'] == before
    assert values['instrumental_results'] == current
    assert values['laboratory_results'] == ''
    assert not issues
    assert origins['provided_documents'] == Path('primary.docx')


def test_ambiguous_examination_dates_keep_block_and_warning():
    text = 'ЭКГ от 09.08.2026 в сравнении с 12.08.2026\nRESULT'
    values, _, issues = partition_examinations(
        [ExaminationSection('instrumental_results', text, Path('mis.docx'))],
        datetime(2026, 8, 10), datetime(2026, 8, 17))
    assert values['instrumental_results'] == text
    assert any(issue.code == 'examination_date_uncertain' for issue in issues)


def test_default_potential_has_no_unrelated_primary_provenance(tmp_path):
    primary = tmp_path / 'primary.docx'
    _write_document(primary, (*_primary_lines(full_name=NAME), 'Реабилитационный потенциал: высокий'))
    draft = scan_discharge_summary(tmp_path)
    assert draft.rehabilitation_potential == 'средний'
    assert 'rehabilitation_potential' not in draft.field_sources


@pytest.mark.parametrize('with_mis', [False, True])
def test_primary_examinations_survive_with_or_without_mis(tmp_path, with_mis):
    primary = tmp_path / 'primary.docx'
    _write_document(primary, (*_primary_lines(full_name=NAME),
        'Лабораторные исследования:', 'ОАК', '09.08.2026', 'BEFORE LAB',
        'Инструментальные исследования:', 'ЭКГ', '10.08.2026', 'CURRENT ECG'))
    if with_mis:
        _write_document(tmp_path / 'mis.docx', _discharge_lines(full_name=NAME))
    draft = scan_discharge_summary(tmp_path)
    assert 'ОАК\n09.08.2026\nBEFORE LAB' in draft.provided_documents
    assert 'ЭКГ\n10.08.2026\nCURRENT ECG' in draft.instrumental_results
    assert 'BEFORE LAB' not in draft.laboratory_results
    assert draft.field_sources['instrumental_results'] == primary
