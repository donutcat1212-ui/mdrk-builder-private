from copy import deepcopy
from dataclasses import replace
from datetime import timedelta

import pytest
from docx import Document

from mdrk_builder.application.admission_only_scales import without_admission_scale_items
from mdrk_builder.application.snapshot import build_snapshot
from mdrk_builder.domain.document_dates import end_of_day, final_mdrk_datetime
from mdrk_builder.application.validation import current_issues
from mdrk_builder.domain import (
    DischargeScaleRow, DischargeSummaryDraft, DischargeTeamFinding,
    MdrkKind, ScaleMeasurement, SpecialistRole,
)
from mdrk_builder.infrastructure.discharge_summary_writer import write_discharge_summary_docx
from mdrk_builder.infrastructure.docx_writer import write_mdrk_docx
from mdrk_builder.infrastructure.draft_store import encode, decode
from test_daily_workflow_ui import app
from test_docx_writer import _representative_episode


def _text(path):
    return '\n'.join(Document(path).element.xpath('//w:t/text()'))


@pytest.mark.parametrize('kind', [MdrkKind.INITIAL, MdrkKind.FINAL])
def test_manual_meeting_changes_header_without_reselecting_clinical_data(app, kind):
    app.episode = _representative_episode(app.episode.folder)
    app._scan_baseline = deepcopy(app.episode)
    app._current_kind = kind
    app.kind_var.set(kind.value)
    app._populate_from_episode()
    before = build_snapshot(app.episode, kind)
    original_output = write_mdrk_docx(app.episode, kind, app.episode.folder / 'original-date.docx', ignore_issues=True)
    original_signatures = next(table for table in Document(original_output).tables
                               if table.cell(0, 0).text == 'Специалист МДРК')
    original = app.episode.meeting_at(kind)
    original_boundary = app.episode.assessment_at(kind)
    changed = original - timedelta(days=2, hours=3)
    if kind is MdrkKind.FINAL:
        changed = final_mdrk_datetime(changed)
    app._entry_variables['meeting'].set(changed.strftime('%d.%m.%Y %H:%M'))
    assert app._apply_form()
    assert not app._test_errors
    assert app.episode.meeting_at(kind) == changed
    assert app.episode.assessment_at(kind) == original_boundary
    after = build_snapshot(app.episode, kind)
    assert replace(after, meeting_at=before.meeting_at) == before
    assert not any(i.code in {'meeting_before_admission', 'final_meeting_not_after_initial',
                             'physician_source_after_meeting'} for i in current_issues(app.episode, kind))
    assert any(label.startswith('Ручная правка') for label, _ in app._field_source_links('meeting_at'))
    app._pending_manual_state = app._capture_manual_state()
    fresh = deepcopy(app._scan_baseline)
    app._merge_manual_state(fresh)
    assert fresh.meeting_at(kind) == changed and fresh.assessment_at(kind) == original_boundary
    restored = decode(encode(fresh))
    assert restored.meeting_at(kind) == changed and restored.assessment_at(kind) == original_boundary
    assert build_snapshot(restored, kind) == after
    output = write_mdrk_docx(restored, kind, app.episode.folder / 'manual-date.docx', ignore_issues=True)
    expected_heading = ('"03" июня 2026 г. время: 13 час. 00 мин.' if kind is MdrkKind.INITIAL
                        else '"17" июня 2026 г. время: 11 час. 00 мин.')
    assert expected_heading in _text(output)
    revised_signatures = next(table for table in Document(output).tables
                              if table.cell(0, 0).text == 'Специалист МДРК')
    assert [[cell.text for cell in row.cells] for row in revised_signatures.rows] == [
        [cell.text for cell in row.cells] for row in original_signatures.rows]


def test_legacy_episode_keeps_its_original_selection_boundary(tmp_path):
    episode = _representative_episode(tmp_path)
    saved = encode(episode)
    saved['fields'].pop('assessment_meetings')
    restored = decode(saved)
    assert restored.assessment_at(MdrkKind.INITIAL) == episode.initial_meeting_at
    assert restored.assessment_at(MdrkKind.FINAL) == end_of_day(episode.final_meeting_at)


def test_egfr_is_only_rendered_in_mdrk1_and_sources_are_unchanged(tmp_path):
    episode = _representative_episode(tmp_path)
    role = SpecialistRole.NEUROLOGIST
    for finding in episode.findings:
        if finding.role is role:
            finding.scales.append(ScaleMeasurement('СКФ', '57,22', finding.source_datetime, role, finding.source))
    for sections in (episode.initial_sections, episode.sections):
        sections.laboratory_results = 'ДРУГОЙ ПОКАЗАТЕЛЬ\nСКФ: 57,22\nЕЩЁ ОДИН ПОКАЗАТЕЛЬ'
    baseline = deepcopy(episode)
    first = write_mdrk_docx(episode, MdrkKind.INITIAL, tmp_path / 'initial.docx', ignore_issues=True)
    final = write_mdrk_docx(episode, MdrkKind.FINAL, tmp_path / 'final.docx', ignore_issues=True)
    assert 'СКФ' in _text(first) and '57,22' in _text(first)
    assert 'СКФ' not in _text(final) and '57,22' not in _text(final)
    assert 'ДРУГОЙ ПОКАЗАТЕЛЬ' in _text(final) and 'ЕЩЁ ОДИН ПОКАЗАТЕЛЬ' in _text(final)
    assert episode == baseline


def test_discharge_omits_egfr_in_all_tables_and_text_including_legacy_drafts(app):
    role = SpecialistRole.NEUROLOGIST
    rows = (DischargeScaleRow(role=role, name='СКФ', value='57,22'),
            DischargeScaleRow(role=role, name='Бартел', value='80'))
    draft = DischargeSummaryDraft(app.episode.folder, admission_scale_rows=rows, discharge_scale_rows=rows,
        team_findings=(DischargeTeamFinding(role=role, conclusion='ЗАКЛЮЧЕНИЕ\nСКФ: 57,22', scales=rows),),
        neurological_status='СТАТУС\nСКФ: 57,22\nСОХРАНИТЬ', discharge_condition='eGFR: 57,22\nСОСТОЯНИЕ')
    original = deepcopy(draft)
    output = write_discharge_summary_docx(draft, app.episode.folder / 'discharge.docx', ignore_issues=True)
    text = _text(output)
    assert 'СКФ' not in text and 'eGFR' not in text and '57,22' not in text
    assert 'Бартел' in text and 'СТАТУС' in text and 'СОХРАНИТЬ' in text
    assert draft == original
    panel = app.discharge_workspace
    panel.load(draft)
    assert all(row.name == 'Бартел' for row in (*panel.draft.admission_scale_rows,
               *panel.draft.discharge_scale_rows, *panel.draft.team_findings[0].scales))
    assert 'СКФ' not in panel._widgets['neurological_status'].get('1.0', 'end-1c')
    state = panel.capture_state()
    state.text['neurological_status'] = original.neurological_status
    panel.restore_state(state)
    assert 'СКФ' not in panel._widgets['neurological_status'].get('1.0', 'end-1c')


def test_removing_egfr_item_preserves_neighbouring_semicolon_items():
    assert without_admission_scale_items('ШРМ: 3; СКФ: 57,22; Бартел: 80') == 'ШРМ: 3; Бартел: 80'


def test_egfr_mention_within_clinical_diagnosis_preserves_complete_diagnosis():
    diagnosis = 'Хроническая болезнь почек 3А ст СКФ 57,22 (МКБ10 N18.3)'
    assert without_admission_scale_items(diagnosis) == diagnosis


def test_egfr_sentence_inside_anthropometry_keeps_neighbouring_facts():
    text = ('Антропометрия: рост 170 см. Клиренс креатинина: 90 мл/мин (Кокрофт-Голт). '
            'Скорость клубочковой фильтрации: 80,5 мл/мин/1,73кв.м (MDRD). '
            'Питание сохранено.')
    assert without_admission_scale_items(text) == (
        'Антропометрия: рост 170 см. Клиренс креатинина: 90 мл/мин (Кокрофт-Голт). '
        'Питание сохранено.')
