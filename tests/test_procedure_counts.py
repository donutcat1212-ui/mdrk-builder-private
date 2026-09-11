from copy import deepcopy
from datetime import datetime
import hashlib

import pytest
from docx import Document

from mdrk_builder.application.discharge_summary import scan_discharge_summary
from mdrk_builder.application.procedures import select_procedures
from mdrk_builder.application.scanner import scan_patient_folder
from mdrk_builder.application.snapshot import build_snapshot
from mdrk_builder.application.validation import current_issues
from mdrk_builder.domain import Episode, MdrkKind, Procedure
from mdrk_builder.infrastructure.discharge_summary_writer import write_discharge_summary_docx
from mdrk_builder.infrastructure.docx_writer import write_mdrk_docx
from mdrk_builder.infrastructure.draft_store import decode, encode
from test_discharge_summary import _primary_lines, _discharge_lines, _write_document
from test_document_order import SYNTHETIC_NAME
from test_docx_writer import _find_table


@pytest.mark.parametrize("planned,actual,manual,expected", [
    (None, 4, False, 4), (10, 4, False, 10), (0, 4, False, 0),
    (None, 0, False, 0), (None, None, False, None), (None, 4, True, None),
])
def test_initial_count_uses_marks_unless_a_plan_or_manual_blank_exists(tmp_path, planned, actual, manual, expected):
    procedure = Procedure("ЛФК", "ФТ", actual, planned_count=planned, source=tmp_path / "assignments.docx")
    if manual:
        procedure.manual_fields.add("planned_count")
    original = deepcopy(procedure)
    selected, = select_procedures([procedure], None, None, MdrkKind.INITIAL)
    assert selected.actual_count == expected
    assert procedure == original
    episode = Episode(tmp_path, procedures=[procedure])
    notices = [issue for issue in current_issues(episode, MdrkKind.INITIAL)
               if issue.code == "planned_count_from_executions"]
    assert bool(notices) == (planned is None and expected is not None and not manual)
    if notices:
        assert notices[0].source == procedure.source
    assert not any(issue.code == "planned_count_from_executions"
                   for issue in current_issues(episode, MdrkKind.FINAL))


@pytest.mark.parametrize("planned_count", [None, 10])
def test_assignment_marks_reach_all_three_documents_and_keep_source_unchanged(tmp_path, planned_count):
    _write_document(tmp_path / "primary.docx", (*_primary_lines(full_name=SYNTHETIC_NAME), "Дата осмотра: 10.08.2026 11:00"))
    _write_document(tmp_path / "mis.docx", _discharge_lines(full_name=SYNTHETIC_NAME))
    source = tmp_path / "assignments.docx"
    document = Document()
    document.add_paragraph("Лист назначений консилиума")
    document.add_paragraph(f"ФИО пациента: {SYNTHETIC_NAME}")
    document.add_paragraph("Номер ИБ: СКП5906/26")
    headers = ["Назначения", "Время", "Кабинет"]
    plan_cells = []
    if planned_count is not None:
        headers.append("Назначено")
        plan_cells.append(str(planned_count))
    headers += ["10.08.2026", "11.08.2026", "17.08.2026", "18.08.2026", "Примечания"]
    table = document.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in [
        ["A19.23.002.014 Индивидуальное занятие ЛФК", "10:55", "1", *plan_cells, "10.55+", "++", "15.30+", "+", "30 мин +"],
        ["A13.23.007 Медико-логопедическая процедура", "11:00", "2", *([""] if plan_cells else []), "", "", "", "", "20 мин"],
    ]:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    document.save(source)
    hashes = {path: hashlib.sha256(path.read_bytes()).digest() for path in tmp_path.glob("*.docx")}
    episode = scan_patient_folder(tmp_path, initial_meeting_at=datetime(2026, 8, 11, 8),
                                  final_meeting_at=datetime(2026, 8, 17, 11))
    assert [row.actual_count for row in episode.procedures] == [4, 0]
    assert episode.procedures[0].planned_count == planned_count
    assert episode.procedures[0].count_needs_review
    assert episode.procedures[0].source_paths == (source,)
    initial_counts = [planned_count if planned_count is not None else 4, 0]
    for candidate in (episode, decode(encode(episode))):
        for kind, expected in [(MdrkKind.INITIAL, initial_counts), (MdrkKind.FINAL, [3, 0])]:
            assert [row.actual_count for row in build_snapshot(candidate, kind).procedures] == expected
    draft = scan_discharge_summary(tmp_path, discharge_datetime_override=datetime(2026, 8, 17, 12))
    assert [row.actual_count for row in draft.completed_procedures] == [3, 0]
    outputs = tmp_path / "output"
    outputs.mkdir()
    for kind, expected in [(MdrkKind.INITIAL, initial_counts), (MdrkKind.FINAL, [3, 0])]:
        output = Document(write_mdrk_docx(episode, kind, outputs / f"{kind.value}.docx", ignore_issues=True))
        rows = _find_table(output, "Реабилитационные процедуры").rows[1:]
        assert [row.cells[2].text for row in rows] == list(map(str, expected))
    output = Document(write_discharge_summary_docx(draft, outputs / "discharge.docx", ignore_issues=True))
    rows = _find_table(output, "Реабилитационные процедуры").rows[1:]
    assert [row.cells[2].text for row in rows] == ["3", "0"]
    assert all(hashlib.sha256(path.read_bytes()).digest() == digest for path, digest in hashes.items())
